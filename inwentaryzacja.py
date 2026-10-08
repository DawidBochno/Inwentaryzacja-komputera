#!/usr/bin/env python3
"""Inwentaryzacja komputera z Windows: sprzet (producent, model, numer seryjny,
procesor, RAM, dyski ze stanem SMART, monitory), Windows i aktywacja, ostatnia
aktualizacja, siec (IP/MAC), drukarki, zainstalowane programy.
Tylko odczyt - nic nie zmienia w systemie.

Wynik: karta komputera XLSX + wiersz w zbiorczym zestawienie.xlsx
(arkusz Komputery: wiersz na komputer, arkusz Programy: programy wszystkich komputerow).

Uruchomienie: python inwentaryzacja.py                       (GUI)
              python inwentaryzacja.py --bez-okna [folder]   (bez okna, np. skrypt logowania)
              python inwentaryzacja.py --selftest            (test)
"""
import base64
import datetime
import json
import os
import subprocess
import sys
import threading
from collections import Counter

if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

MAX_DNI_AKTUALIZACJI = 35  # jak w Kontroli stanowiska
OBUDOWY = {3: "stacjonarny", 4: "stacjonarny", 5: "stacjonarny", 6: "stacjonarny (mini tower)",
           7: "stacjonarny (tower)", 8: "laptop", 9: "laptop", 10: "laptop", 13: "all-in-one",
           14: "laptop", 15: "stacjonarny", 16: "stacjonarny", 17: "serwer", 23: "serwer (rack)",
           24: "stacjonarny", 30: "tablet", 31: "laptop 2w1", 32: "laptop 2w1", 35: "mini PC", 36: "stacjonarny"}
PAMIEC = {20: "DDR", 21: "DDR2", 24: "DDR3", 26: "DDR4", 29: "LPDDR3", 30: "LPDDR4", 34: "DDR5", 35: "LPDDR5"}
LICENCJA = {0: "brak licencji", 1: "aktywowany", 2: "okres prolongaty", 3: "okres prolongaty",
            4: "okres prolongaty (nieoryginalny)", 5: "nieaktywowany (powiadomienia)", 6: "przedłużony okres prolongaty"}
STAN_DYSKU = {"Healthy": "dobry", "Warning": "ostrzeżenie", "Unhealthy": "zły"}

# Zbieranie faktow: PowerShell 5.1 (jest w kazdym Windows 10/11), bez uprawnien
# administratora. Kazdy odczyt w try - brak danych = null, nie blad calosci.
SKRYPT_PS = r"""
$ErrorActionPreference = 'Stop'
$ProgressPreference = 'SilentlyContinue'
[Console]::OutputEncoding = [Text.Encoding]::UTF8
function Proba($b) { try { & $b } catch { $null } }
function Czas($d) { if ($d) { ([datetime]$d).ToString('s') } else { $null } }
function Rejestr($p, $n) { try { (Get-ItemProperty -Path $p -Name $n).$n } catch { $null } }
function Cim($k) { try { Get-CimInstance $k | Select-Object -First 1 } catch { $null } }
function Tekst($a) { if ($a) { (-join ($a | Where-Object { $_ } | ForEach-Object { [char]$_ })).Trim() } }
$id = [Security.Principal.WindowsIdentity]::GetCurrent()
$adm = ([Security.Principal.WindowsPrincipal]$id).IsInRole([Security.Principal.WindowsBuiltInRole]::Administrator)
$os = Cim Win32_OperatingSystem
$cs = Cim Win32_ComputerSystem
$bios = Cim Win32_BIOS
$pl = Cim Win32_BaseBoard
$uninst = 'HKLM:\SOFTWARE\Microsoft\Windows\CurrentVersion\Uninstall\*',
  'HKLM:\SOFTWARE\WOW6432Node\Microsoft\Windows\CurrentVersion\Uninstall\*',
  'HKCU:\Software\Microsoft\Windows\CurrentVersion\Uninstall\*'
$f = [ordered]@{
  komputer = $env:COMPUTERNAME
  uzytkownik = "$env:USERDOMAIN\$env:USERNAME"
  admin = $adm
  w_domenie = [bool]$cs.PartOfDomain
  domena = $cs.Domain
  producent = $cs.Manufacturer
  model = $cs.Model
  ram = [double]$cs.TotalPhysicalMemory
  obudowa = Proba { [int]@((Cim Win32_SystemEnclosure).ChassisTypes)[0] }
  numer = "$($bios.SerialNumber)".Trim()
  bios = "$($bios.SMBIOSBIOSVersion)"
  bios_data = Czas $bios.ReleaseDate
  plyta = "$($pl.Manufacturer) $($pl.Product)".Trim()
  procesory = Proba { Get-CimInstance Win32_Processor | ForEach-Object { [ordered]@{
    nazwa = "$($_.Name)".Trim(); rdzenie = $_.NumberOfCores; watki = $_.NumberOfLogicalProcessors } } }
  pamiec = Proba { Get-CimInstance Win32_PhysicalMemory | ForEach-Object { [ordered]@{
    rozmiar = [double]$_.Capacity; typ = [int]$_.SMBIOSMemoryType; mhz = $_.ConfiguredClockSpeed } } }
  gniazda = Proba { (Get-CimInstance Win32_PhysicalMemoryArray | Measure-Object MemoryDevices -Sum).Sum }
  grafika = Proba { @(Get-CimInstance Win32_VideoController | ForEach-Object { $_.Name }) }
  monitory = Proba { Get-CimInstance -Namespace root/wmi -ClassName WmiMonitorID | ForEach-Object { [ordered]@{
    producent = Tekst $_.ManufacturerName; model = Tekst $_.UserFriendlyName; numer = Tekst $_.SerialNumberID;
    rok = $_.YearOfManufacture } } }
  dyski = Proba { Get-PhysicalDisk | ForEach-Object { $d = $_; [ordered]@{ nazwa = "$($d.FriendlyName)";
    typ = "$($d.MediaType)"; magistrala = "$($d.BusType)"; rozmiar = [double]$d.Size;
    numer = "$($d.SerialNumber)".Trim(); stan = "$($d.HealthStatus)";
    smart = if ($adm) { Proba { $r = $d | Get-StorageReliabilityCounter; [ordered]@{
      temperatura = $r.Temperature; zuzycie = $r.Wear; godziny = $r.PowerOnHours } } } } } }
  woluminy = Proba { Get-CimInstance Win32_LogicalDisk -Filter 'DriveType=3' | ForEach-Object { [ordered]@{
    dysk = $_.DeviceID; etykieta = $_.VolumeName; system_plikow = $_.FileSystem;
    wolne = [double]$_.FreeSpace; rozmiar = [double]$_.Size } } }
  system = $os.Caption
  kompilacja = $os.BuildNumber
  wersja = Rejestr 'HKLM:\SOFTWARE\Microsoft\Windows NT\CurrentVersion' 'DisplayVersion'
  architektura = $os.OSArchitecture
  zainstalowany = Czas $os.InstallDate
  uruchomiony = Czas $os.LastBootUpTime
  aktywacja = Proba { Get-CimInstance SoftwareLicensingProduct -Filter ("ApplicationID='55c92734-d682-4d71-983e-d6ec3f16059f'" +
    " AND PartialProductKey IS NOT NULL") | ForEach-Object { [ordered]@{
    opis = $_.Description; stan = [int]$_.LicenseStatus; klucz = $_.PartialProductKey } } }
  aktualizacja = Czas (Proba { (New-Object -ComObject Microsoft.Update.AutoUpdate).Results.LastInstallationSuccessDate })
  siec = Proba { Get-CimInstance Win32_NetworkAdapterConfiguration -Filter 'IPEnabled=TRUE' | ForEach-Object { [ordered]@{
    karta = $_.Description; mac = $_.MACAddress; ip = @($_.IPAddress); brama = @($_.DefaultIPGateway);
    dns = @($_.DNSServerSearchOrder); dhcp = [bool]$_.DHCPEnabled } } }
  drukarki = Proba { Get-CimInstance Win32_Printer | ForEach-Object { [ordered]@{ nazwa = $_.Name;
    sterownik = $_.DriverName; port = $_.PortName; domyslna = [bool]$_.Default; sieciowa = [bool]$_.Network } } }
  programy = Proba { Get-ItemProperty $uninst -ErrorAction SilentlyContinue |
    Where-Object { $_.DisplayName -and $_.SystemComponent -ne 1 -and -not $_.ParentKeyName } | ForEach-Object {
    [ordered]@{ nazwa = "$($_.DisplayName)".Trim(); wersja = "$($_.DisplayVersion)"; wydawca = "$($_.Publisher)";
    data = "$($_.InstallDate)" } } }
}
$f | ConvertTo-Json -Depth 6 -Compress
"""

# ---------------------------------------------------------- zbieranie ----


def zbierz():
    """Uruchamia SKRYPT_PS i zwraca fakty (slownik)."""
    kod = base64.b64encode(SKRYPT_PS.encode("utf-16-le")).decode()
    r = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
         "-EncodedCommand", kod],
        capture_output=True, timeout=300, creationflags=0x08000000)  # CREATE_NO_WINDOW
    wyjscie = r.stdout.decode("utf-8", "replace").lstrip("﻿").strip()
    if not wyjscie.startswith("{"):
        raise RuntimeError("PowerShell nie zwrocil danych: %s" % (
            r.stderr.decode("utf-8", "replace").strip()[:300] or wyjscie[:300]))
    return json.loads(wyjscie.splitlines()[-1])

# ------------------------------------------------------------- analiza ----


def lista(v):
    """ConvertTo-Json zwraca obiekt zamiast listy przy 1 elemencie, null przy 0."""
    if v is None:
        return []
    return [x for x in v if x is not None] if isinstance(v, list) else [v]


def data(s):
    return datetime.datetime.fromisoformat(s) if s else None


def dzien(s):
    d = data(s)
    return d.strftime("%Y-%m-%d") if d else ""


def data_instalacji(s):
    """InstallDate z rejestru: zwykle RRRRMMDD, ale instalatory wpisuja rozne rzeczy."""
    try:
        return datetime.datetime.strptime(s, "%Y%m%d").strftime("%Y-%m-%d")
    except (TypeError, ValueError):
        return s or ""


def gb(b, dziel=1e9):
    return "%.0f GB" % (b / dziel) if b else ""


def analizuj(f, teraz):
    """Fakty -> (lista arkuszy [(nazwa, naglowki, wiersze)], lista uwag).
    Pierwszy arkusz "Komputer" (pole, wartosc) - zawsze te same pola w tej samej
    kolejnosci, bo to kolumny zbiorczego zestawienia."""
    uwagi = []

    # pamiec RAM
    pam = lista(f.get("pamiec"))
    ram = sum(p.get("rozmiar") or 0 for p in pam) or f.get("ram") or 0
    opis_ram = gb(ram, 2 ** 30)
    if pam:
        kosci = Counter(round((p.get("rozmiar") or 0) / 2 ** 30) for p in pam)
        opis_ram += ": " + " + ".join("%d x %d GB" % (n, r) for r, n in sorted(kosci.items(), reverse=True))
        typ, mhz = PAMIEC.get(pam[0].get("typ")), pam[0].get("mhz")
        opis_ram += "".join(" " + str(x) for x in (typ, mhz and "%s MHz" % mhz) if x)
        if f.get("gniazda"):
            opis_ram += ", gniazda %d/%d" % (len(pam), f["gniazda"])

    # dyski
    dyski = []
    for d in lista(f.get("dyski")):
        s = d.get("smart") or {}
        stan = STAN_DYSKU.get(d.get("stan"), d.get("stan") or "")
        if d.get("stan") and d["stan"] != "Healthy":
            uwagi.append("dysk %s: stan %s" % (d.get("nazwa"), stan))
        dyski.append([d.get("nazwa"), "" if d.get("typ") == "Unspecified" else d.get("typ"), d.get("magistrala"),
                      gb(d.get("rozmiar")), d.get("numer"), stan, s.get("temperatura"), s.get("zuzycie"),
                      s.get("godziny")])
    woluminy = []
    for w in lista(f.get("woluminy")):
        proc = 100 * w["wolne"] / w["rozmiar"] if w.get("rozmiar") else None
        if proc is not None and proc < 10:
            uwagi.append("mało miejsca na %s (%.0f%%)" % (w["dysk"], proc))
        woluminy.append([w.get("dysk"), w.get("etykieta"), w.get("system_plikow"), gb(w.get("rozmiar")),
                         gb(w.get("wolne")), "%.0f%%" % proc if proc is not None else ""])

    # aktywacja Windows: produkt z kluczem; aktywowany ma pierwszenstwo
    akt = sorted(lista(f.get("aktywacja")), key=lambda a: a.get("stan") != 1)
    if akt:
        a = akt[0]
        kanal = (a.get("opis") or "").rsplit(",", 1)[-1].replace("channel", "").strip()
        opis_akt = "%s (%s, klucz ...%s)" % (LICENCJA.get(a.get("stan"), "stan %s" % a.get("stan")),
                                             kanal, a.get("klucz"))
        if a.get("stan") != 1:
            uwagi.append("Windows nieaktywowany")
    else:
        opis_akt = "nie udało się odczytać"

    aktualizacja = data(f.get("aktualizacja"))
    if aktualizacja is None or (teraz - aktualizacja).days > MAX_DNI_AKTUALIZACJI:
        uwagi.append("brak aktualizacji Windows od ponad %d dni" % MAX_DNI_AKTUALIZACJI)

    siec = [[n.get("karta"), n.get("mac"), ", ".join(lista(n.get("ip"))), ", ".join(lista(n.get("brama"))),
             ", ".join(lista(n.get("dns"))), "tak" if n.get("dhcp") else "nie"] for n in lista(f.get("siec"))]
    ipv4 = [ip for n in lista(f.get("siec")) for ip in lista(n.get("ip")) if ":" not in ip]
    drukarki = [[d.get("nazwa"), d.get("sterownik"), d.get("port"), "tak" if d.get("domyslna") else "",
                 "tak" if d.get("sieciowa") else ""] for d in lista(f.get("drukarki"))]
    monitory = [[m.get("producent"), m.get("model"), m.get("numer"), m.get("rok")] for m in lista(f.get("monitory"))]
    programy = sorted({(p.get("nazwa"), p.get("wersja"), p.get("wydawca"), data_instalacji(p.get("data")))
                       for p in lista(f.get("programy"))}, key=lambda p: (p[0] or "").lower())

    typ = f.get("obudowa")
    kompilacja = "%s %s (kompilacja %s, %s)" % (f.get("system") or "?", f.get("wersja") or "",
                                                f.get("kompilacja"), f.get("architektura") or "")
    up = data(f.get("uruchomiony"))
    komputer = [
        ("Komputer", f.get("komputer")),
        ("Data inwentaryzacji", teraz.strftime("%Y-%m-%d %H:%M")),
        ("Użytkownik", f.get("uzytkownik")),
        ("Domena", f.get("domena") if f.get("w_domenie") else "poza domeną (%s)" % (f.get("domena") or "")),
        ("Typ", OBUDOWY.get(typ, "inny (%s)" % typ if typ else "")),
        ("Producent", f.get("producent")),
        ("Model", f.get("model")),
        ("Numer seryjny", f.get("numer")),
        ("Płyta główna", f.get("plyta")),
        ("BIOS", "%s (%s)" % (f.get("bios") or "", dzien(f.get("bios_data")))),
        ("Procesor", "; ".join("%s (%s rdzeni, %s wątków)" % (p.get("nazwa"), p.get("rdzenie"), p.get("watki"))
                               for p in lista(f.get("procesory")))),
        ("Pamięć RAM", opis_ram),
        ("Karta graficzna", "; ".join(lista(f.get("grafika")))),
        ("Dyski", "; ".join(" ".join(str(x) for x in (d[1], d[3], d[2] and "(%s)" % d[2]) if x) for d in dyski)),
        ("Monitory", "; ".join("%s %s (SN %s)" % (m[0], m[1], m[2]) for m in monitory)),
        ("System", kompilacja),
        ("Zainstalowano system", dzien(f.get("zainstalowany"))),
        ("Aktywacja Windows", opis_akt),
        ("Ostatnia aktualizacja", dzien(f.get("aktualizacja")) or "brak danych"),
        ("Ostatnie uruchomienie", up.strftime("%Y-%m-%d %H:%M") if up else ""),
        ("Adres IP", ", ".join(ipv4)),
        ("Adres MAC", ", ".join(n[1] for n in siec if n[1])),
        ("Drukarki", len(drukarki)),
        ("Programy", len(programy)),
        ("Uwagi", "; ".join(uwagi)),
    ]
    arkusze = [
        ("Komputer", ["Pole", "Wartość"], [list(k) for k in komputer]),
        ("Programy", ["Program", "Wersja", "Wydawca", "Data instalacji"], [list(p) for p in programy]),
        ("Dyski", ["Dysk", "Typ", "Magistrala", "Rozmiar", "Numer seryjny", "Stan (SMART)",
                   "Temperatura °C", "Zużycie %", "Godziny pracy"], dyski),
        ("Woluminy", ["Wolumin", "Etykieta", "System plików", "Rozmiar", "Wolne", "Wolne %"], woluminy),
        ("Sieć", ["Karta", "MAC", "Adres IP", "Brama", "DNS", "DHCP"], siec),
        ("Drukarki", ["Drukarka", "Sterownik", "Port", "Domyślna", "Sieciowa"], drukarki),
        ("Monitory", ["Producent", "Model", "Numer seryjny", "Rok produkcji"], monitory),
    ]
    return arkusze, uwagi

# ------------------------------------------------------------- zapis ----


def wypelnij(ws, naglowki, wiersze):
    from openpyxl.styles import Font
    ws.append(naglowki)
    for w in wiersze:
        ws.append(["" if v is None else v for v in w])
    for c in ws[1]:
        c.font = Font(bold=True)
        szer = max(len(str(ws.cell(r, c.column).value or "")) for r in range(1, min(ws.max_row, 200) + 1))
        ws.column_dimensions[c.column_letter].width = min(max(szer + 2, 10), 70)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def zapisz_karte(arkusze, sciezka):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nazwa, nag, wiersze in arkusze:
        wypelnij(wb.create_sheet(nazwa), nag, wiersze)
    wb.save(sciezka)


def zapisz_zestawienie(arkusze, sciezka):
    """Arkusz Komputery: wiersz na komputer; Programy: programy wszystkich komputerow.
    Ponowna inwentaryzacja tego samego komputera zastepuje jego wiersze."""
    import openpyxl
    from openpyxl.styles import PatternFill
    pola = arkusze[0][2]
    komputer = pola[0][1]
    nowe = [("Komputery", [p for p, _ in pola], [[v for _, v in pola]]),
            ("Programy", ["Komputer"] + arkusze[1][1], [[komputer] + p for p in arkusze[1][2]])]
    if os.path.exists(sciezka):
        wb = openpyxl.load_workbook(sciezka)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    for nazwa, nag, wiersze in nowe:
        stare = []
        if nazwa in wb.sheetnames:
            ws = wb[nazwa]
            stary_nag = [c.value for c in ws[1]]
            # po kolumnach z naglowka - nowa wersja programu moze dodac pole
            stare = [[dict(zip(stary_nag, r)).get(n) for n in nag]
                     for r in ws.iter_rows(min_row=2, values_only=True) if r[0] != komputer]
            wb.remove(ws)
        ws = wb.create_sheet(nazwa)
        wypelnij(ws, nag, sorted(stare + wiersze, key=lambda w: str(w[0]).lower()))
    ws = wb["Komputery"]
    wb.move_sheet(ws, -wb.index(ws))
    kol = [c.value for c in ws[1]].index("Uwagi") + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, kol).value:
            ws.cell(r, kol).fill = PatternFill("solid", fgColor="FFC7CE")
    wb.save(sciezka)


def run(out_dir, log, zrodlo=None):
    """Inwentaryzacja, karta XLSX, wiersz w zestawieniu. Zwraca (arkusze, uwagi, plik_karty)."""
    log("Odczyt danych komputera (do ok. 30 s) ...")
    f = (zrodlo or zbierz)()
    teraz = datetime.datetime.now()
    arkusze, uwagi = analizuj(f, teraz)
    log("%s: %d programow, %d dyskow, %d drukarek, uwag: %d" % (
        f.get("komputer"), len(arkusze[1][2]), len(arkusze[2][2]), len(arkusze[5][2]), len(uwagi)))
    os.makedirs(out_dir, exist_ok=True)
    plik = os.path.join(out_dir, "inwentaryzacja_%s_%s.xlsx" % (f.get("komputer"), teraz.strftime("%Y-%m-%d")))
    zapisz_karte(arkusze, plik)
    log("Karta komputera: " + plik)
    zest = os.path.join(out_dir, "zestawienie.xlsx")
    try:
        zapisz_zestawienie(arkusze, zest)
        log("Zestawienie: " + zest)
    except PermissionError:
        log("BLAD: zestawienie.xlsx jest otwarte (np. w Excelu) - nie dopisano wiersza. Karta zapisana.")
    return arkusze, uwagi, plik

# ---------------------------------------------------------------- GUI ----


def gui():
    import ctypes
    import tkinter as tk
    from tkinter import scrolledtext, ttk

    root = tk.Tk()
    root.title("Inwentaryzacja komputera")
    root.geometry("1100x660")
    pad = dict(padx=6, pady=3)
    v_out = tk.StringVar(value=os.path.join(APP_DIR, "OUTPUT"))
    stan = {"plik": None}

    f = ttk.Frame(root)
    f.pack(fill="x", **pad)
    ttk.Label(f, text="Folder wyjsciowy:").pack(side="left", **pad)
    ttk.Entry(f, textvariable=v_out, width=70).pack(side="left", **pad)
    admin = bool(ctypes.windll.shell32.IsUserAnAdmin())
    ttk.Label(root, text="Uprawnienia: administrator" if admin else
              "Uprawnienia: zwykly uzytkownik (wystarcza; jako administrator dodatkowo temperatura "
              "i zuzycie dyskow)", foreground="gray").pack(anchor="w", padx=12)

    h = ttk.Frame(root)
    h.pack(fill="x", **pad)
    btn = ttk.Button(h, text="Inwentaryzuj")
    btn.pack(side="left", padx=6)
    b_karta = ttk.Button(h, text="Otworz karte", state="disabled", command=lambda: os.startfile(stan["plik"]))
    b_karta.pack(side="left", padx=6)
    ttk.Button(h, text="Otworz zestawienie", command=lambda: (
        os.startfile(os.path.join(v_out.get(), "zestawienie.xlsx"))
        if os.path.exists(os.path.join(v_out.get(), "zestawienie.xlsx"))
        else log("Zestawienie jeszcze nie istnieje - kliknij Inwentaryzuj."))).pack(side="left", padx=6)

    def jako_admin():
        skrypt = os.path.abspath(__file__)
        if ctypes.windll.shell32.ShellExecuteW(None, "runas", sys.executable, '"%s"' % skrypt, APP_DIR, 1) > 32:
            root.destroy()

    if not admin:
        ttk.Button(h, text="Uruchom jako administrator", command=jako_admin).pack(side="left", padx=6)

    zakladki = ttk.Notebook(root)
    zakladki.pack(fill="both", expand=True, **pad)
    log_box = scrolledtext.ScrolledText(root, height=5)
    log_box.pack(fill="x", **pad)

    def log(msg):
        def put():
            log_box.insert("end", str(msg) + "\n")
            log_box.see("end")
        root.after(0, put)

    def pokaz(arkusze, plik):
        for z in zakladki.tabs():
            zakladki.forget(z)
        for nazwa, nag, wiersze in arkusze:
            ramka = ttk.Frame(zakladki)
            t = ttk.Treeview(ramka, columns=nag, show="headings")
            for k in nag:
                t.heading(k, text=k)
                t.column(k, width=800 if nag == ["Pole", "Wartość"] and k == "Wartość" else 160, anchor="w")
            t.tag_configure("UWAGA", background="#FFD7D7")
            for w in wiersze:
                t.insert("", "end", values=["" if v is None else v for v in w],
                         tags=("UWAGA",) if w[0] == "Uwagi" and w[1] else ())
            sb = ttk.Scrollbar(ramka, orient="vertical", command=t.yview)
            t.configure(yscrollcommand=sb.set)
            sb.pack(side="right", fill="y")
            t.pack(side="left", fill="both", expand=True)
            zakladki.add(ramka, text="%s (%d)" % (nazwa, len(wiersze)) if nazwa != "Komputer" else nazwa)
        stan["plik"] = plik
        b_karta.config(state="normal")

    def start():
        log_box.delete("1.0", "end")
        btn.config(state="disabled")

        def work():
            try:
                arkusze, _, plik = run(v_out.get().strip('" '), log)
                root.after(0, lambda: pokaz(arkusze, plik))
            except Exception as e:
                log("BLAD: %s" % e)
            finally:
                root.after(0, lambda: btn.config(state="normal"))

        threading.Thread(target=work, daemon=True).start()

    btn.config(command=start)
    if "--selftest" in sys.argv:
        root.after(200, root.destroy)
    import aktualizacja
    aktualizacja.start(root, "DawidBochno/Inwentaryzacja-komputera", "main", "inwentaryzacja.py")
    root.mainloop()

# ------------------------------------------------------------- selftest ----


def selftest():
    import tempfile
    import openpyxl

    teraz = datetime.datetime(2026, 10, 8, 12, 0)
    iso = lambda dni: (teraz - datetime.timedelta(days=dni)).strftime("%Y-%m-%dT%H:%M:%S")
    pc = {
        "komputer": "UG-PC-001", "uzytkownik": "URZAD\\jkowalska", "admin": True,
        "w_domenie": True, "domena": "urzad.local", "producent": "Dell Inc.", "model": "OptiPlex 7090",
        "ram": 17000000000.0, "obudowa": 3, "numer": "ABC1234", "bios": "1.21.0", "bios_data": "2024-03-01T00:00:00",
        "plyta": "Dell Inc. 0XYZ12",
        "procesory": {"nazwa": "Intel(R) Core(TM) i5-10500 CPU @ 3.10GHz", "rdzenie": 6, "watki": 12},
        "pamiec": [{"rozmiar": 8 * 2 ** 30, "typ": 26, "mhz": 2666}, {"rozmiar": 8 * 2 ** 30, "typ": 26, "mhz": 2666}],
        "gniazda": 4, "grafika": ["Intel(R) UHD Graphics 630"],
        "monitory": [{"producent": "DEL", "model": "DELL P2419H", "numer": "CN0ABC", "rok": 2021}],
        "dyski": [{"nazwa": "KXG60ZNV512G", "typ": "SSD", "magistrala": "NVMe", "rozmiar": 512e9, "numer": "X1",
                   "stan": "Healthy", "smart": {"temperatura": 38, "zuzycie": 4, "godziny": 9120}},
                  {"nazwa": "USB Flash", "typ": "Unspecified", "magistrala": "USB", "rozmiar": 32e9, "numer": "",
                   "stan": "Healthy", "smart": None}],
        "woluminy": {"dysk": "C:", "etykieta": "", "system_plikow": "NTFS", "wolne": 200e9, "rozmiar": 510e9},
        "system": "Microsoft Windows 11 Pro", "kompilacja": "26200", "wersja": "25H2", "architektura": "64-bitowy",
        "zainstalowany": "2025-01-10T09:00:00", "uruchomiony": iso(1),
        "aktywacja": {"opis": "Windows(R) Operating System, OEM_DMI channel", "stan": 1, "klucz": "3V66T"},
        "aktualizacja": iso(5),
        "siec": {"karta": "Intel(R) Ethernet I219-LM", "mac": "AA:BB:CC:00:11:22",
                 "ip": ["192.168.1.15", "fe80::1"], "brama": ["192.168.1.1"], "dns": None, "dhcp": True},
        "drukarki": [{"nazwa": "HP LaserJet M404", "sterownik": "HP Universal", "port": "IP_192.168.1.50",
                      "domyslna": True, "sieciowa": False},
                     {"nazwa": "Microsoft Print to PDF", "sterownik": "x", "port": "PORTPROMPT:",
                      "domyslna": False, "sieciowa": False}],
        "programy": [{"nazwa": "7-Zip 24.08", "wersja": "24.08", "wydawca": "Igor Pavlov", "data": ""},
                     {"nazwa": "Adobe Acrobat Reader", "wersja": "24.1", "wydawca": "Adobe", "data": "20240115"},
                     {"nazwa": "Adobe Acrobat Reader", "wersja": "24.1", "wydawca": "Adobe", "data": "20240115"}],
    }
    arkusze, uwagi = analizuj(pc, teraz)
    k = dict(arkusze[0][2])
    assert uwagi == [] and k["Uwagi"] == "", uwagi
    assert k["Typ"] == "stacjonarny" and k["Numer seryjny"] == "ABC1234"
    assert k["Pamięć RAM"] == "16 GB: 2 x 8 GB DDR4 2666 MHz, gniazda 2/4", k["Pamięć RAM"]
    assert k["Procesor"] == "Intel(R) Core(TM) i5-10500 CPU @ 3.10GHz (6 rdzeni, 12 wątków)"
    assert k["Dyski"] == "SSD 512 GB (NVMe); 32 GB (USB)", k["Dyski"]
    assert k["Aktywacja Windows"] == "aktywowany (OEM_DMI, klucz ...3V66T)", k["Aktywacja Windows"]
    assert k["Adres IP"] == "192.168.1.15" and k["Adres MAC"] == "AA:BB:CC:00:11:22"
    assert k["Monitory"] == "DEL DELL P2419H (SN CN0ABC)" and k["Drukarki"] == 2 and k["Programy"] == 2
    assert k["BIOS"] == "1.21.0 (2024-03-01)" and k["Domena"] == "urzad.local"
    ark = {a[0]: a for a in arkusze}
    assert ark["Programy"][2] == [["7-Zip 24.08", "24.08", "Igor Pavlov", ""],
                                  ["Adobe Acrobat Reader", "24.1", "Adobe", "2024-01-15"]]
    assert ark["Dyski"][2][0][5:] == ["dobry", 38, 4, 9120] and ark["Woluminy"][2][0][5] == "39%"
    assert ark["Sieć"][2][0][2:] == ["192.168.1.15, fe80::1", "192.168.1.1", "", "tak"]

    zly = dict(pc, komputer="UG-PC-002", aktualizacja=None, obudowa=10, pamiec=None, gniazda=None,
               aktywacja=[{"opis": "Windows(R) Operating System, RETAIL channel", "stan": 5, "klucz": "AAAAA"}],
               dyski={"nazwa": "ST1000", "typ": "HDD", "magistrala": "SATA", "rozmiar": 1e12, "numer": "Z1",
                      "stan": "Warning", "smart": None},
               woluminy=[{"dysk": "C:", "wolne": 5e9, "rozmiar": 100e9}], siec=None, drukarki=None,
               programy={"nazwa": "LibreOffice", "wersja": "24.2", "wydawca": "TDF", "data": "2024"})
    arkusze2, uwagi = analizuj(zly, teraz)
    k = dict(arkusze2[0][2])
    assert uwagi == ["dysk ST1000: stan ostrzeżenie", "mało miejsca na C: (5%)", "Windows nieaktywowany",
                     "brak aktualizacji Windows od ponad 35 dni"], uwagi
    assert k["Typ"] == "laptop" and k["Pamięć RAM"] == "16 GB" and k["Ostatnia aktualizacja"] == "brak danych"
    assert k["Aktywacja Windows"].startswith("nieaktywowany (powiadomienia) (RETAIL")
    assert k["Adres IP"] == "" and k["Programy"] == 1
    k = dict(analizuj(dict(pc, aktywacja=None, w_domenie=False, domena="WORKGROUP", obudowa=99), teraz)[0][0][2])
    assert k["Aktywacja Windows"] == "nie udało się odczytać" and k["Typ"] == "inny (99)"
    assert k["Domena"] == "poza domeną (WORKGROUP)"
    assert [a[0] for a in arkusze] == [a[0] for a in arkusze2]
    assert [p for p, _ in arkusze[0][2]] == [p for p, _ in arkusze2[0][2]]  # stale kolumny zestawienia
    assert lista(None) == [] and lista({"a": 1}) == [{"a": 1}] and lista([None, 1]) == [1]
    assert data_instalacji("2024") == "2024" and data_instalacji(None) == ""

    # zapis: karta + zestawienie (zastapienie wierszy, kolumny po naglowkach)
    tmp = tempfile.mkdtemp()
    lines = []
    _, _, plik = run(tmp, lines.append, zrodlo=lambda: pc)
    assert all(ord(ch) < 128 for line in lines for ch in line), lines
    wb = openpyxl.load_workbook(plik)
    assert wb.sheetnames == ["Komputer", "Programy", "Dyski", "Woluminy", "Sieć", "Drukarki", "Monitory"]
    assert wb["Komputer"]["B9"].value == "ABC1234" and wb["Programy"].max_row == 3
    run(tmp, lines.append, zrodlo=lambda: zly)
    run(tmp, lines.append, zrodlo=lambda: dict(pc, numer="NOWY1"))  # ponownie PC-001
    zest = os.path.join(tmp, "zestawienie.xlsx")
    wb = openpyxl.load_workbook(zest)
    assert wb.sheetnames == ["Komputery", "Programy"], wb.sheetnames
    ws = wb["Komputery"]
    nag = [c.value for c in ws[1]]
    assert [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)] == ["UG-PC-001", "UG-PC-002"]
    assert ws.cell(2, nag.index("Numer seryjny") + 1).value == "NOWY1"
    uw = ws.cell(3, nag.index("Uwagi") + 1)
    assert "nieaktywowany" in uw.value and uw.fill.fgColor.rgb.endswith("FFC7CE")
    assert [r[:2] for r in wb["Programy"].iter_rows(min_row=2, values_only=True)] == [
        ("UG-PC-001", "7-Zip 24.08"), ("UG-PC-001", "Adobe Acrobat Reader"), ("UG-PC-002", "LibreOffice")]
    # starsze zestawienie bez kolumny (np. po dodaniu pola w nowej wersji) - wiersze po naglowkach
    ws.delete_cols(nag.index("Monitory") + 1)
    wb.save(zest)
    run(tmp, lines.append, zrodlo=lambda: dict(zly, komputer="UG-PC-003"))
    ws = openpyxl.load_workbook(zest)["Komputery"]
    assert [c.value for c in ws[1]] == nag and ws.cell(2, nag.index("Numer seryjny") + 1).value == "NOWY1"
    assert ws.cell(2, nag.index("Monitory") + 1).value is None and ws.max_row == 4

    # prawdziwy odczyt z tego komputera (w CI: czysty Windows Server)
    if sys.platform == "win32":
        f = zbierz()
        for klucz in ("komputer", "system", "kompilacja", "procesory", "woluminy", "siec", "programy"):
            assert f.get(klucz), (klucz, f.get(klucz))
        arkusze, uwagi = analizuj(f, datetime.datetime.now())
        k = dict(arkusze[0][2])
        assert k["Procesor"] and k["Programy"] > 0 and k["Adres IP"], k
        print("ten komputer: %s %s, %s programow, uwagi: %s" % (
            k["Producent"], k["Model"], k["Programy"], k["Uwagi"] or "brak"))

    import aktualizacja
    aktualizacja.selftest()
    print("selftest OK")


def bez_okna():
    """--bez-okna [folder]: inwentaryzacja bez okna (np. skrypt logowania, zadanie)."""
    i = sys.argv.index("--bez-okna")
    out = sys.argv[i + 1] if len(sys.argv) > i + 1 else os.path.join(APP_DIR, "OUTPUT")
    run(out, print)


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
        if "--gui" in sys.argv:
            gui()
    elif "--bez-okna" in sys.argv:
        bez_okna()
    else:
        gui()
